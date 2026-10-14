import pickle,re,zipfile,collections
from lxml import etree
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
rows=pickle.load(open('rows.pkl','rb'))
wb=openpyxl.Workbook()
hdrfill=PatternFill('solid',fgColor='1F3864'); hf=Font(name='BIZ UDPゴシック',bold=True,color='FFFFFF'); nf=Font(name='BIZ UDPゴシック')
def sheet(ws,head,data,widths):
    ws.append(head)
    for r in data: ws.append(r)
    for c in ws[1]: c.fill=hdrfill;c.font=hf
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font=nf; c.alignment=Alignment(wrap_text=True,vertical='top')
    for i,w in enumerate(widths): ws.column_dimensions[chr(65+i)].width=w
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
ws=wb.active; ws.title='入力分類'
sheet(ws,['シート','セル','元の値','分類(案)','テンプレートでの扱い','根拠(近くのラベル等)','関連質問'],rows,[16,8,34,22,22,40,9])
Q=[('Q1','「使い方」シートを追加するか','既定=テンプレート版のみ(整形版は元の構成のまま)'),
('Q2','配色','紺ベース/緑ベース/今の水色基調を活かす。おすすめ=紺ベース＋グレー(入力=薄い水色)'),
('Q3','シート保護','数式セルのみロック・パスワードなし/あり。おすすめ=パスワードなし'),
('Q4','Bとして残す値の範囲(事業検討の費用・建築シナジー)','物件王の加盟金(450万)・会費(20万/月)・研修費・PPC/SVコンサル、建築シナジーの単価/利益率/受注頻度、歩留り(5%/30%/50%)、CVR等を「前提」として残す案でよいか'),
('Q5','事業検討!C43 採用費/年(168)','想定年収×35%の概算。A(空)かB(残す)か'),
('Q6','店舗内装の必要面積(J10:J15、M19ほか)','物件王の標準レイアウト値のため B(残す)案。会社ごとに変える値か'),
('Q7','免許シートの費用(D9, D14:D21)','登録免許税・協会入会金など。都道府県で異なる可能性。B(残す)案か、A(空)か'),
('Q8','簡易PL!D7:D9, I7 の件数(2,1,1,6)','数式ではなく手入力。会社固有のためA(空)案。PLの件数と連動していないことも要注意(気づいた点参照)'),
('Q9','目標設定 D38/K38/U38 の■(対応事業チェック)','テンプレート初期値は □(未選択)でよいか'),
('Q10','PLの入力欄(サンプルでも0)','0のまま残す案(空にしても結果は同じ)。空にしたい場合は指示ください'),
('Q11','研修詳細の水色セル','カリキュラム本文(物件王の研修内容)のため C(そのまま)案。会社ごとに変える内容ではない前提でよいか'),
('Q12','PLで水色の数式セル(約548個)を数式用の見た目に統一してよいか','一覧は変更レポートに出します'),
('Q13','決定事項(署名書類)の様式','既定=見た目の微調整のみ・文面は変えない。A4で1ページに収める'),
('Q14','提出・印刷の主な用途','Excel共有/PDF提出/紙印刷。優先順位に影響'),
('Q15','用紙サイズ','現状 全体フロー・研修詳細・店舗内装・事業検討・簡易PL・PL2枚はA3指定。ご指示は「すべてA4」。A3→A4で縮小して収めてよいか(全体フローは特に縮小大)')]
sheet(wb.create_sheet('質問リスト'),['番号','質問','内容・おすすめ'],Q,[7,48,90])
# hardcoded
z=zipfile.ZipFile('orig/orig.xlsx'); ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
names=["社名","全体フロー","研修詳細","店舗内装","ターゲットエリア","免許","目標設定","事業検討","簡易PL","PL（1年目）","PL（2年目以降）","決定事項"]
grp=collections.OrderedDict()
for i,n in enumerate(names,1):
    t=etree.fromstring(z.read(f'xl/worksheets/sheet{i}.xml'))
    for c in t.xpath('//m:sheetData//m:c[m:f]',namespaces=ns):
        ft=c.find('m:f',ns).text or ''
        if not ft: continue
        ft2=re.sub(r"'[^']*'!|[^\s=(,+\-*/<>&]+!",'',ft); ft2=re.sub(r'\$?[A-Z]{1,3}\$?\d+','',ft2)
        nums=[x for x in re.findall(r'\d+\.?\d*%?',ft2) if x not in('0','1')]
        if not nums: continue
        norm=re.sub(r'(\$?[A-Z]{1,3}\$?)\d+',r'\1#',ft); norm=re.sub(r'\$?[A-Z]{1,3}(?=#)','~',norm)
        g=grp.setdefault((n,norm),[ft,[]]); g[1].append(c.get('r'))
out=[]
for (n,norm),(ft,cells) in grp.items():
    out.append([n,', '.join(cells[:6])+(f' …(計{len(cells)}セル)' if len(cells)>6 else ''),'='+ft,len(cells)])
sheet(wb.create_sheet('数式内の直書き定数'),['シート','セル','数式(代表例)','セル数'],out,[16,40,70,8])
print(len(out),sum(r[3] for r in out))
I=[('全体','用紙サイズ','全体フロー・研修詳細・店舗内装・事業検討・簡易PL・PL(1年目)・PL(2年目以降)はA3(paperSize=8)。ご指示のA4とズレ。→Q15'),
('目標設定/事業検討','手数料/件の食い違い','目標設定!J16=ターゲットエリア!AV6*10000(432,932円)、事業検討!E4=ROUNDDOWN(…,-3)(432,000円)。同じ「手数料/件」が画面上2つの値。数式は変更せず報告'),
('事業検討','O43','=IF(L51=5000000,0,L43)。加盟金L51は4,500,000のため条件は常に不成立(500万を想定した書き方)。意図要確認'),
('事業検討','L54(店舗改装費用)','入力欄だが空欄。C49(減価償却)・L56・L57・L58に影響。サンプルでは未入力'),
('事業検討','W36','=X37*Z3 。Z3は空のため結果0。残骸の可能性'),
('簡易PL','D7:D9, I7','件数が手入力(2/1/1/6)。PLの件数(D6=PL!R3等)とは非連動。I8,I9は事業検討に連動で不統一'),
('PL(2年目以降)','D6,E6,G6… など','同じ行で数式セルと定数0が混在(例: 行6)。入力の意図か要確認'),
('PL(1年目/2年目以降)','水色の数式セル','水色が入力の色なのに数式が約548個。色だけでは入力と判別不可(見た目統一の対象)'),
('印刷範囲','事業検討','データはAB列まで、印刷範囲はA:M。作業領域(O〜AB)を意図的に外している可能性大→変更せず維持案'),
('印刷範囲','簡易PL/PL','簡易PLはA:K(作業領域AL〜BI除外)、PLはA:R(W〜AI除外)。意図的な外し方と判断し維持案'),
('印刷範囲','全体フロー/ターゲットエリア/店舗内装/社名','印刷範囲未設定(全体フローはGO列×1003行の書式セルが約19.7万個)'),
('免許','D10, D22','=SUM(D2:E9) 等、範囲がD9より広く(D3の文字を含む)。結果は正しいが範囲が緩い'),
('数値の文字列化','全シート','数値が文字列で入っているセルは0件'),
('ターゲットエリア','AO6','配列数式 =SUM(Z12:AN12*110%)(動的配列メタデータ cm=1)。保持必須')]
sheet(wb.create_sheet('気づいた点'),['対象','場所','内容(数式は直していません)'],I,[18,30,100])
import os; os.makedirs('/home/user/-/deliverables',exist_ok=True)
wb.save('/home/user/-/deliverables/入力分類表_案.xlsx')
