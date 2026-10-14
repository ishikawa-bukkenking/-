import sys,re,json,csv,pickle,collections,zipfile
sys.path.insert(0,'/tmp/claude-0/-home-user--/db782ef1-e569-59dc-97e4-a77616c95c04/scratchpad/tools')
from make_template import *
from openpyxl.utils import get_column_letter as gl
R='/home/user/-/deliverables/report/'
t=json.load(open(SD+'out/template.json')); o=json.load(open(SD+'orig/recalc.json'))
# ---- 1. エラー一覧と起点 ----
err={(sh,c) for sh,d in t.items() for c,v in d.items() if v['e']!=0}
oerr={(sh,c) for sh,d in o.items() for c,v in d.items() if v['e']!=0}
ref_re=re.compile(r"(?:\$?(?:'([^']+)'|([^.\s;()=+\-*/<>&$'!:,{}]+))\.)?\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")
def refs(sh,f):
    out=set()
    for m in ref_re.finditer(f):
        s_=m.group(1) or m.group(2) or sh
        c1,r1=ci(m.group(3)),int(m.group(4)); c2,r2=(ci(m.group(5)),int(m.group(6))) if m.group(5) else (c1,r1)
        if (c2-c1+1)*(r2-r1+1)>3000: continue
        for r in range(r1,r2+1):
            for c in range(c1,c2+1): out.add((s_,f'{gl(c)}{r}'))
    return out
rows=[]; roots=[]
for sh,c in sorted(err,key=lambda x:(NAMES.index(x[0]) if x[0] in NAMES else 99,)+tuple(re.match(r'([A-Z]+)(\d+)',x[1]).groups()[::-1])):
    f=t[sh][c]['f']; rf=refs(sh,f)-{(sh,c)}
    is_root=not any(x in err for x in rf)
    if is_root: roots.append((sh,c))
    rows.append([sh,c,t[sh][c]['s'],lo_to_excel(f),'起点' if is_root else '波及','元のサンプルでもエラー' if (sh,c) in oerr else ''])
with open(R+'テンプレート_エラー一覧.csv','w',newline='',encoding='utf-8-sig') as fp:
    w=csv.writer(fp); w.writerow(['シート','セル','表示','数式','種別(起点=自分の式が空入力で0割りなど / 波及=他のエラーを参照)','備考']); w.writerows(rows)
print('エラー',len(err),'起点',len(roots)); print(roots)
by=collections.Counter(r[0] for r in rows); print(by)
pickle.dump((rows,roots),open(SD+'out/errrows.pkl','wb'))
# ---- 2. PLの水色の数式セル ----
zo=zipfile.ZipFile(SRC)
st=etree.fromstring(zo.read('xl/styles.xml')); xfs=st.xpath('//m:cellXfs/m:xf',namespaces=N); fills=st.xpath('//m:fills/m:fill',namespaces=N)
def is_cyan(s):
    fe=fills[int(xfs[s].get('fillId'))].find('m:patternFill/m:fgColor',N); return fe is not None and fe.get('rgb')=='FFCCFFFF'
rows2=[]
for i in(10,11):
    name=NAMES[i-1]; root=etree.fromstring(zo.read(f'xl/worksheets/sheet{i}.xml'))
    for c in root.iterfind('.//m:sheetData/m:row/m:c',N):
        if c.find('m:f',N) is not None and is_cyan(int(c.get('s','0'))):
            rc=o[name].get(c.get('r'),{}); rows2.append([name,c.get('r'),lo_to_excel(rc.get('f','')),rc.get('s','')])
with open(R+'PL_水色の数式セル一覧.csv','w',newline='',encoding='utf-8-sig') as fp:
    w=csv.writer(fp); w.writerow(['シート','セル','数式(元の水色=入力の色が付いていた数式セル)','計算値(表示)']); w.writerows(rows2)
print('PL水色の数式',len(rows2),collections.Counter(r[0] for r in rows2))
# ---- 3. 直書き定数 ----
cl=constants_list()
with open(R+'数式内の直書き定数一覧.csv','w',newline='',encoding='utf-8-sig') as fp:
    w=csv.writer(fp); w.writerow(['シート','セル(例)','数式(例)','該当セル数']); w.writerows(cl)
print('定数パターン',len(cl),sum(x[3] for x in cl))
# ---- 4. 入力分類(確定) ----
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
rows=pickle.load(open(SD+'rows.pkl','rb')); out=[]
roles=load_roles()
keys={(r[0],r[1]) for r in rows}
for sh,c,val,k,act,why,q in rows:
    cls=k; act2=act
    if sh=='事業検討' and c=='C43': cls='A 会社固有→空にする'; act2='空にする'; why='採用費/年(ご回答Q5: A)'
    if sh=='簡易PL' and c in('D7','D8','D9','I7'): cls='B 想定値→残す'; act2='残す'; why='件数(ご回答Q8: 残す)'
    if sh in('PL（1年目）','PL（2年目以降）') and k.startswith('A'): act2='0のまま(ご回答Q10)'
    if (sh,c) in CHECKS: act2='□(未選択)にする(ご回答Q9)'
    if (sh,c)==('社名','C15'): act2='「株式会社〇〇」(仮の社名)に置き換える(整形版・テンプレート版とも)'; why='社名(ご指示で仮名に変更)'
    if k.startswith('A') and act2=='空にする' : pass
    out.append([sh,c,val,cls,act2,why])
wb=openpyxl.Workbook(); ws=wb.active; ws.title='入力分類(確定)'
ws.append(['シート','セル','元の値','分類','テンプレート版での扱い','根拠']); 
for r in out: ws.append(r)
for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor=PAL['BRAND'])
for k,w in zip('ABCDEF',(16,8,34,22,28,46)): ws.column_dimensions[k].width=w
ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
wb.save('/home/user/-/deliverables/report/入力分類表_確定.xlsx')
print(collections.Counter(r[3] for r in out))
